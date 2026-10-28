# -*- coding: utf-8 -*-
"""
ایمپورت دسته‌جمعی با تطبیق هوشمند نام مراکز.
- اگر نام دقیقاً نبود، نزدیک‌ترین مرکز انتخاب می‌شود.
- مراکز ناشناخته در خروجی گزارش می‌شوند تا کاربر دستی تخصیص دهد.
"""
import sys
import os
import difflib
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from db.database import get_connection, log_action

# ==== CONTRACT_KEYWORD_MAP از importer اصلی ====
try:
    from importer.excel_importer import CONTRACT_KEYWORD_MAP
except Exception:
    CONTRACT_KEYWORD_MAP = []

_DIGIT_MAP = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789"
)


def _to_number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return 0.0
    s = s.translate(_DIGIT_MAP).replace(",", "").replace("،", "").replace(" ", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _normalize_text(value):
    if value is None:
        return None
    if not isinstance(value, str):
        return value
    v = value.strip()
    v = v.replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
    v = " ".join(v.split())
    return v if v else None


# ============================================================
# امتیازدهی ستون‌ها
# ============================================================
def _score_center_header(h):
    if not h:
        return -100
    h = str(h).replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
    score = 0
    if "عنوان" in h: score -= 50
    if "شماره" in h or "کد" in h: score -= 30
    if "تاریخ" in h: score -= 30
    if "ردیف" == h.strip(): score -= 50
    if "مرکز هزینه" in h: score += 30
    elif "مرکز" in h: score += 20
    elif "شبکه" in h: score += 15
    elif "واحد" in h: score += 5
    return score


def _score_amount_header(h):
    if not h:
        return -100
    h = str(h).replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
    score = 0
    if "عنوان" in h: score -= 50
    if "ردیف" == h.strip(): score -= 50
    if "توضیحات" in h: score -= 50
    if "تاریخ" in h: score -= 50
    if "شماره" in h: score -= 30
    if "تعداد" in h: score -= 20
    if "واحد" in h: score -= 15
    if "سال" in h: score -= 30
    if "ماه" in h: score -= 30

    if "هزینه کل" in h or "مبلغ کل" in h or "جمع کل" in h or "مبلغ نهایی" in h:
        score += 100
    elif "هزینه" in h and "واحد" not in h:
        score += 40
    elif "مبلغ" in h:
        score += 60
    elif "جمع" in h:
        score += 50
    elif "کارشناس" in h or "تجهیزات" in h or "جایزه" in h or "هدایا" in h or "پاداش" in h:
        score += 40
    elif "ریال" in h or "تومان" in h:
        score += 30
    return score


def _score_center_data(col_idx, data_rows):
    score = 0
    for row in data_rows:
        if col_idx >= len(row):
            continue
        v = row[col_idx]
        if v is None:
            continue
        s = str(v).strip()
        if not s:
            continue
        cleaned = s.translate(_DIGIT_MAP).replace(",", "").replace(" ", "")
        if cleaned.replace(".", "").isdigit():
            score -= 5
            continue
        if 2 <= len(s) <= 80:
            score += 10
    return score


def _score_amount_data(col_idx, data_rows):
    score = 0
    for row in data_rows:
        if col_idx >= len(row):
            continue
        v = row[col_idx]
        if v is None:
            continue
        if isinstance(v, (int, float)) and v > 1000:
            score += 20
            continue
        s = str(v).strip()
        if not s:
            continue
        cleaned = s.translate(_DIGIT_MAP).replace(",", "").replace("،", "").replace(" ", "")
        cleaned = cleaned.replace(".", "")
        if cleaned.isdigit() and int(cleaned) > 1000:
            score += 20
    return score


def _detect_columns(ws, header_row_idx, data_start_row):
    try:
        header_row = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))
    except StopIteration:
        header_row = []

    data_rows = []
    try:
        for i, row in enumerate(ws.iter_rows(min_row=data_start_row, max_row=data_start_row + 9, values_only=True)):
            data_rows.append(row)
            if i >= 9:
                break
    except Exception:
        pass

    num_cols = max(
        len(header_row),
        max((len(r) for r in data_rows if r), default=0)
    )

    center_scores = []
    amount_scores = []
    for i in range(num_cols):
        h = header_row[i] if i < len(header_row) else None
        center_scores.append(_score_center_header(h) * 10 + _score_center_data(i, data_rows))
        amount_scores.append(_score_amount_header(h) * 10 + _score_amount_data(i, data_rows))

    best_center = max(range(num_cols), key=lambda i: center_scores[i]) if num_cols else 0
    best_amount = max(range(num_cols), key=lambda i: amount_scores[i]) if num_cols else 1

    if best_center == best_amount and num_cols > 1:
        sorted_a = sorted(range(num_cols), key=lambda i: amount_scores[i], reverse=True)
        for i in sorted_a:
            if i != best_center:
                best_amount = i
                break

    return best_center, best_amount


def _scan_file(ws, max_rows=15):
    for ri in range(1, max_rows + 1):
        try:
            row = next(ws.iter_rows(min_row=ri, max_row=ri, values_only=True))
        except StopIteration:
            break
        text = " ".join(str(v) for v in row if v is not None)
        markers = ["مرکز هزینه", "مرکز", "مبلغ", "هزینه کل", "جمع", "عنوان"]
        found = sum(1 for m in markers if m in text)
        if found >= 2:
            try:
                next_row = next(ws.iter_rows(min_row=ri + 1, max_row=ri + 1, values_only=True))
                if next_row and any(v is not None for v in next_row):
                    return ri, ri + 1
            except StopIteration:
                pass
    return None, 1


def _read_file(xlsx_path, max_scan_rows=15):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row, data_start_row = _scan_file(ws, max_scan_rows)

    if header_row:
        col_center, col_amount = _detect_columns(ws, header_row, data_start_row)
        mode = "header"
    else:
        data_rows = []
        try:
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
                data_rows.append(row)
                if i >= 9:
                    break
        except Exception:
            pass

        num_cols = max((len(r) for r in data_rows if r), default=0)
        center_scores = [_score_center_data(i, data_rows) for i in range(num_cols)]
        amount_scores = [_score_amount_data(i, data_rows) for i in range(num_cols)]

        col_center = max(range(num_cols), key=lambda i: center_scores[i]) if num_cols else 0
        col_amount = max(range(num_cols), key=lambda i: amount_scores[i]) if num_cols else 1
        if col_center == col_amount and num_cols > 1:
            sorted_a = sorted(range(num_cols), key=lambda i: amount_scores[i], reverse=True)
            for i in sorted_a:
                if i != col_center:
                    col_amount = i
                    break
        data_start_row = 1
        mode = "auto"

    sums = {}
    total_rows = 0
    total_amount = 0
    sample_rows = []

    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        if not row:
            continue
        total_rows += 1

        def get(idx):
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        center_raw = get(col_center)
        amount_raw = get(col_amount)
        center = _normalize_text(center_raw)
        amount = _to_number(amount_raw)

        if len(sample_rows) < 5:
            sample_rows.append({
                "center": str(center_raw)[:60] if center_raw else "",
                "amount": str(amount_raw)[:30] if amount_raw is not None else "",
            })

        if not center or amount <= 0:
            continue

        sums[center] = sums.get(center, 0) + amount
        total_amount += amount

    wb.close()

    debug = {
        "mode": mode,
        "header_row": header_row,
        "start_row": data_start_row,
        "col_center": col_center,
        "col_amount": col_amount,
        "total_rows": total_rows,
        "unique_centers": len(sums),
        "total_amount": total_amount,
        "sample_rows": sample_rows,
    }

    return sums, total_rows, debug


# ============================================================
# تطبیق هوشمند نام مراکز (Fuzzy)
# ============================================================
def _find_best_match(name, known_centers):
    """
    پیدا کردن نزدیک‌ترین مرکز موجود برای نام داده‌شده.
    returns (matched_name, method) یا (None, None)
    """
    if not name:
        return None, None

    # ۱) تطابق دقیق
    if name in known_centers:
        return name, "exact"

    # ۲) تطابق از CONTRACT_KEYWORD_MAP
    for keyword, target in CONTRACT_KEYWORD_MAP:
        if keyword in name:
            if target in known_centers:
                return target, f"keyword"
            # شاید target با normalized فرق کنه
            for k in known_centers:
                if target in k or k in target:
                    return k, f"keyword"

    # ۳) تطابق substring — اگه یه مرکز داخل اسم باشه
    for k in sorted(known_centers.keys(), key=lambda x: -len(x)):
        if len(k) >= 4 and k in name:
            return k, "substring"

    # ۴) تطابق substring معکوس
    for k in sorted(known_centers.keys(), key=lambda x: -len(x)):
        if len(name) >= 4 and name in k:
            return k, "substring_rev"

    # ۵) fuzzy match
    matches = difflib.get_close_matches(
        name, list(known_centers.keys()), n=1, cutoff=0.55
    )
    if matches:
        return matches[0], "fuzzy"

    return None, None


def _get_all_known_centers(conn):
    """همه مراکز شناخته‌شده از centers_structure"""
    known = {}
    for r in conn.execute("""
        SELECT DISTINCT center_name, supply_type, deputy
        FROM centers_structure
    """).fetchall():
        cname = _normalize_text(r["center_name"])
        if cname:
            known[cname] = {
                "supply_type": r["supply_type"] or "",
                "deputy": r["deputy"] or "",
            }
    return known


# ============================================================
# تابع اصلی
# ============================================================
def import_bulk(xlsx_path, category, db_path, username="سیستم", period_month=None):
    if category not in ("specialists", "equipment", "bonus"):
        raise ValueError("دستهٔ نامعتبر")
    if not period_month:
        raise ValueError("period_month الزامی است")

    sums, total_rows, debug = _read_file(xlsx_path)

    if not sums:
        msg = "هیچ دادهٔ معتبری در فایل پیدا نشد.\n"
        msg += f"ردیف‌های اسکن‌شده: {total_rows}\n"
        if debug['sample_rows']:
            msg += "نمونهٔ ردیف‌های خوانده‌شده:\n"
            for i, s in enumerate(debug['sample_rows'], 1):
                msg += f"  {i}. مرکز: «{s['center']}» | مبلغ: «{s['amount']}»\n"
        raise ValueError(msg)

    conn = get_connection(db_path)
    known_centers = _get_all_known_centers(conn)

    existing_rows = {}
    for r in conn.execute("""
        SELECT id, center_name, supply_type
        FROM specialists_equipment_bonus
        WHERE period_month = ?
    """, (period_month,)).fetchall():
        key = (_normalize_text(r["center_name"]) or "", r["supply_type"] or "")
        existing_rows[key] = r["id"]

    col_name = {
        "specialists": "specialists",
        "equipment": "equipment",
        "bonus": "bonus",
    }[category]

    inserted = 0
    updated = 0
    remapped = []           # مراکزی که نامشون با تطبیق هوشمند عوض شد
    unmatched = []          # مراکزی که هیچ تطبیقی نداشتن

    for raw_center, amount in sums.items():
        # پیدا کردن مرکز معتبر
        matched, method = _find_best_match(raw_center, known_centers)

        if not matched:
            # بدون تطبیق — با نام اصلی ذخیره می‌کنیم
            unmatched.append({
                "raw_name": raw_center,
                "amount": amount,
            })
            final_center = raw_center
            supply_type = ""
        else:
            final_center = matched
            info = known_centers.get(matched, {})
            supply_type = info.get("supply_type", "")
            if method != "exact":
                remapped.append({
                    "original": raw_center,
                    "matched": matched,
                    "method": method,
                    "amount": amount,
                })

        key = (final_center, supply_type)

        if key in existing_rows:
            row_id = existing_rows[key]
            conn.execute(
                f"UPDATE specialists_equipment_bonus SET {col_name} = ? WHERE id = ?",
                (amount, row_id)
            )
            updated += 1
        else:
            defaults = {"specialists": 0.0, "equipment": 0.0, "bonus": 0.0}
            defaults[col_name] = amount
            try:
                conn.execute("""
                    INSERT INTO specialists_equipment_bonus
                        (center_name, supply_type, period_month, specialists, equipment, bonus)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (final_center, supply_type, period_month,
                      defaults["specialists"],
                      defaults["equipment"],
                      defaults["bonus"]))
                inserted += 1
                # اضافه به existing
                new_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                existing_rows[key] = new_id
                if final_center not in known_centers:
                    known_centers[final_center] = {"supply_type": supply_type, "deputy": ""}
            except Exception as e:
                print(f"[BULK] خطا در درج {final_center}: {e}")

    conn.commit()

    label = {
        "specialists": "کارشناسان",
        "equipment": "خدمات و تجهیزات",
        "bonus": "جوایز و هدایا",
    }[category]

    log_action(conn, username, f"ایمپورت دسته‌جمعی {label}",
               f"{len(sums)} مرکز، درج: {inserted}، به‌روز: {updated}، "
               f"تطبیق هوشمند: {len(remapped)}، ناشناخته: {len(unmatched)}")

    conn.close()

    return {
        "category": category,
        "category_label": label,
        "centers_count": len(sums),
        "total_rows_read": total_rows,
        "inserted": inserted,
        "updated": updated,
        "remapped": remapped,
        "remapped_count": len(remapped),
        "unmatched": unmatched,
        "unmatched_count": len(unmatched),
        "debug": debug,
    }