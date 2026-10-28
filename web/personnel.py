# -*- coding: utf-8 -*-
"""بانک نفرات — دیتابیس + منطق"""
import sqlite3, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
from app_paths import get_db_path


JOB_TITLES = [
    "ارزیاب برنامه",
    "ارزیاب پخش",
    "پژوهشگر",
    "تدوینگر",
    "تصویر بردار رسانه",
    "تهیه کننده تلویزیونی",
    "تهیه کننده رادیویی",
    "چهره پرداز",
    "صدا بردار",
    "طراح دکور",
    "فیلم نامه نویس",
    "کاردان تاسیسات",
    "کاردان صدا و تصویر",
    "کارشناس برنامه‌ریزی محتوایی",
    "کارگر امور حرفه‌ای",
    "کارگردان پویا نمایی",
    "کارگردان تلویزیونی",
    "کارمند امور دفتری",
    "کارمند هماهنگی برنامه",
    "گرافیست",
    "گزارشگر رسانه",
    "گوینده رسانه",
    "مترجم",
    "متصدی صحنه تلویزیونی",
    "مدیر پخش",
    "منشی صحنه تلویزیونی",
    "مهندس انتقال و انتشار",
    "مهندس سخت افزار و شبکه",
    "مهندس صدا و تصویر",
    "مهندس نرم افزار",
    "نورپرداز",
    "نویسنده",
    "هنرپشه",
    "ویراستار",
    "سایر",
]

CONTRACT_TYPES = [
    "رسمی",
    "پیمانی",
    "قرارداد خرید خدمات",
    "شرکتی",
    "ماده 28",
]

SUPPLY_TYPES = ["تامین", "پخش", "تولید"]

SIMA_DEPUTY = "معاونت سیما"


def _conn():
    c = sqlite3.connect(str(get_db_path()), timeout=30)
    c.row_factory = sqlite3.Row
    return c


def init_personnel_tables():
    con = _conn()
    con.executescript("""
        CREATE TABLE IF NOT EXISTS personnel (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            national_code TEXT UNIQUE NOT NULL,
            first_name TEXT DEFAULT '',
            last_name TEXT DEFAULT '',
            job_title TEXT DEFAULT '',
            contract_type TEXT DEFAULT '',
            deputy TEXT DEFAULT '',
            center TEXT DEFAULT '',
            supply_type TEXT DEFAULT '',
            status_changed_date TEXT DEFAULT '',
            created_at TEXT DEFAULT (datetime('now')),
            updated_at TEXT DEFAULT (datetime('now'))
        );
        CREATE INDEX IF NOT EXISTS idx_personnel_nc ON personnel(national_code);
        CREATE INDEX IF NOT EXISTS idx_personnel_deputy ON personnel(deputy);
        CREATE INDEX IF NOT EXISTS idx_personnel_center ON personnel(center);

        CREATE TABLE IF NOT EXISTS personnel_transfers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            personnel_id INTEGER NOT NULL REFERENCES personnel(id) ON DELETE CASCADE,
            from_deputy TEXT DEFAULT '',
            from_center TEXT DEFAULT '',
            to_deputy TEXT DEFAULT '',
            to_center TEXT DEFAULT '',
            transfer_date TEXT NOT NULL,
            transfer_method TEXT DEFAULT '',
            note TEXT DEFAULT '',
            created_at TEXT DEFAULT (datetime('now'))
        );
        CREATE INDEX IF NOT EXISTS idx_transfers_pid ON personnel_transfers(personnel_id);

        CREATE TABLE IF NOT EXISTS personnel_status_changes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            personnel_id INTEGER NOT NULL REFERENCES personnel(id) ON DELETE CASCADE,
            from_contract_type TEXT DEFAULT '',
            to_contract_type TEXT NOT NULL,
            change_date TEXT NOT NULL,
            change_method TEXT DEFAULT '',
            note TEXT DEFAULT '',
            created_at TEXT DEFAULT (datetime('now'))
        );
        CREATE INDEX IF NOT EXISTS idx_status_pid ON personnel_status_changes(personnel_id);
    """)

    # ==== Migration: اضافه کردن supply_type اگه نبود ====
    try:
        cols = [r[1] for r in con.execute("PRAGMA table_info(personnel)").fetchall()]
        if "supply_type" not in cols:
            con.execute("ALTER TABLE personnel ADD COLUMN supply_type TEXT DEFAULT ''")
            print("[MIGRATION] ستون supply_type اضافه شد")
    except Exception as e:
        print("[MIGRATION] خطا:", e)

    con.commit()
    con.close()


# ============================================================
# CRUD
# ============================================================
def list_personnel(deputy=None, center=None, search=None, limit=500, offset=0):
    con = _conn()
    q = "SELECT * FROM personnel WHERE 1=1"
    p = []
    if deputy and deputy != "همه":
        q += " AND deputy = ?"; p.append(deputy)
    if center and center != "همه":
        q += " AND center = ?"; p.append(center)
    if search:
        q += " AND (national_code LIKE ? OR first_name LIKE ? OR last_name LIKE ? OR job_title LIKE ?)"
        s = f"%{search}%"
        p.extend([s, s, s, s])
    q += " ORDER BY last_name, first_name LIMIT ? OFFSET ?"
    p.extend([limit, offset])
    rows = con.execute(q, p).fetchall()
    total = con.execute("SELECT COUNT(*) FROM personnel").fetchone()[0]
    con.close()
    return [dict(r) for r in rows], total


def list_personnel_enriched(deputy=None, center=None, search=None,
                             job_title=None, contract_type=None,
                             limit=1000, offset=0):
    con = _conn()
    q = """
        SELECT p.*,
            (SELECT t.from_deputy FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_from_deputy,
            (SELECT t.from_center FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_from_center,
            (SELECT t.to_deputy FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_to_deputy,
            (SELECT t.to_center FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_to_center,
            (SELECT t.transfer_date FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_date,
            (SELECT t.transfer_method FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_method,
            (SELECT COUNT(*) FROM personnel_transfers t WHERE t.personnel_id = p.id) AS tr_count,
            (SELECT s.from_contract_type FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_from,
            (SELECT s.to_contract_type FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_to,
            (SELECT s.change_date FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_date,
            (SELECT s.change_method FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_method,
            (SELECT COUNT(*) FROM personnel_status_changes s WHERE s.personnel_id = p.id) AS st_count
        FROM personnel p
        WHERE 1=1
    """
    params = []
    if deputy and deputy != "همه":
        q += " AND p.deputy = ?"; params.append(deputy)
    if center and center != "همه":
        q += " AND p.center = ?"; params.append(center)
    if job_title and job_title != "همه":
        q += " AND p.job_title = ?"; params.append(job_title)
    if contract_type and contract_type != "همه":
        q += " AND p.contract_type = ?"; params.append(contract_type)
    if search:
        q += " AND (p.national_code LIKE ? OR p.first_name LIKE ? OR p.last_name LIKE ?)"
        s = f"%{search}%"
        params.extend([s, s, s])
    q += " ORDER BY p.last_name, p.first_name LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    rows = con.execute(q, params).fetchall()
    total = con.execute("SELECT COUNT(*) FROM personnel").fetchone()[0]
    con.close()
    return [dict(r) for r in rows], total


def get_personnel(pid):
    con = _conn()
    row = con.execute("SELECT * FROM personnel WHERE id=?", (pid,)).fetchone()
    con.close()
    return dict(row) if row else None


def create_personnel(data):
    nc = (data.get("national_code") or "").strip().zfill(10)
    if not nc:
        raise ValueError("کد ملی الزامی است")
    con = _conn()
    try:
        cur = con.execute("""
            INSERT INTO personnel
                (national_code, first_name, last_name, job_title, contract_type,
                 deputy, center, supply_type, status_changed_date)
            VALUES (?,?,?,?,?,?,?,?,?)
        """, (nc,
              data.get("first_name", ""),
              data.get("last_name", ""),
              data.get("job_title", ""),
              data.get("contract_type", ""),
              data.get("deputy", ""),
              data.get("center", ""),
              data.get("supply_type", ""),
              data.get("status_changed_date", "")))
        con.commit()
        return cur.lastrowid
    except sqlite3.IntegrityError:
        raise ValueError("این کد ملی قبلاً ثبت شده است")
    finally:
        con.close()


def update_personnel(pid, data):
    fields, params = [], []
    for k in ("first_name", "last_name", "job_title", "contract_type",
              "deputy", "center", "supply_type", "status_changed_date"):
        if k in data:
            fields.append(f"{k}=?")
            params.append(data[k])
    if not fields:
        return
    fields.append("updated_at=datetime('now')")
    params.append(pid)
    con = _conn()
    con.execute(f"UPDATE personnel SET {', '.join(fields)} WHERE id=?", params)
    con.commit()
    con.close()


def delete_personnel(pid):
    con = _conn()
    con.execute("DELETE FROM personnel WHERE id=?", (pid,))
    con.commit()
    con.close()


# ============================================================
# انتقال
# ============================================================
def do_transfer(pid, data):
    to_deputy = (data.get("to_deputy") or "").strip()
    to_center = (data.get("to_center") or "").strip()
    to_supply_type = (data.get("to_supply_type") or "").strip()
    transfer_date = (data.get("transfer_date") or "").strip()

    if not to_center or not transfer_date:
        raise ValueError("مقصد و تاریخ الزامی است")

    # برای معاونت سیما، تامین/پخش الزامی
    if to_deputy == SIMA_DEPUTY and not to_supply_type:
        raise ValueError("برای معاونت سیما، تعیین تامین/پخش/تولید الزامی است")

    # برای غیرسیما، خالی باشه
    if to_deputy != SIMA_DEPUTY:
        to_supply_type = ""

    con = _conn()
    p = con.execute("SELECT deputy, center, COALESCE(supply_type,'') AS supply_type FROM personnel WHERE id=?", (pid,)).fetchone()
    if not p:
        con.close()
        raise ValueError("نفر پیدا نشد")

    from_deputy = p["deputy"] or ""
    from_center = p["center"] or ""
    from_supply = p["supply_type"] or ""

    cur = con.execute("""
        INSERT INTO personnel_transfers
            (personnel_id, from_deputy, from_center, to_deputy, to_center,
             transfer_date, transfer_method, note)
        VALUES (?,?,?,?,?,?,?,?)
    """, (pid, from_deputy, from_center, to_deputy, to_center,
          transfer_date,
          data.get("transfer_method", ""),
          data.get("note", "")))

    # به‌روز deputy + center + supply_type
    con.execute("""
        UPDATE personnel
        SET deputy=?, center=?, supply_type=?, updated_at=datetime('now')
        WHERE id=?
    """, (to_deputy, to_center, to_supply_type, pid))

    con.commit()
    new_id = cur.lastrowid
    con.close()
    return new_id


def list_transfers(pid):
    con = _conn()
    rows = con.execute("""
        SELECT * FROM personnel_transfers
        WHERE personnel_id=?
        ORDER BY transfer_date DESC, id DESC
    """, (pid,)).fetchall()
    con.close()
    return [dict(r) for r in rows]


# ============================================================
# تبدیل وضعیت
# ============================================================
def do_status_change(pid, data):
    to_type = (data.get("to_contract_type") or "").strip()
    change_date = (data.get("change_date") or "").strip()
    if not to_type or not change_date:
        raise ValueError("وضعیت جدید و تاریخ الزامی است")

    con = _conn()
    p = con.execute("SELECT contract_type FROM personnel WHERE id=?", (pid,)).fetchone()
    if not p:
        con.close()
        raise ValueError("نفر پیدا نشد")

    from_type = p["contract_type"] or ""

    cur = con.execute("""
        INSERT INTO personnel_status_changes
            (personnel_id, from_contract_type, to_contract_type,
             change_date, change_method, note)
        VALUES (?,?,?,?,?,?)
    """, (pid, from_type, to_type, change_date,
          data.get("change_method", ""),
          data.get("note", "")))

    con.execute("""
        UPDATE personnel SET contract_type=?, status_changed_date=?, updated_at=datetime('now')
        WHERE id=?
    """, (to_type, change_date, pid))
    con.commit()
    new_id = cur.lastrowid
    con.close()
    return new_id


def list_status_changes(pid):
    con = _conn()
    rows = con.execute("""
        SELECT * FROM personnel_status_changes
        WHERE personnel_id=?
        ORDER BY change_date DESC, id DESC
    """, (pid,)).fetchall()
    con.close()
    return [dict(r) for r in rows]


# ============================================================
# ایمپورت از اکسل
# ============================================================
def import_personnel_excel(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_idx = 1
    header = None
    for ri in range(1, 6):
        try:
            row = next(ws.iter_rows(min_row=ri, max_row=ri, values_only=True))
        except StopIteration:
            break
        if not row:
            continue
        text = " ".join(str(v) for v in row if v is not None)
        if "کد ملی" in text or "کدملی" in text:
            header_idx = ri
            header = row
            break

    if header is None:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    col = {"nc": None, "first": None, "last": None, "job": None,
           "contract": None, "deputy": None, "center": None, "supply": None}

    for i, h in enumerate(header):
        if h is None:
            continue
        h_norm = str(h).strip().replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")

        if col["nc"] is None and ("کد ملی" in h_norm or "کدملی" in h_norm):
            col["nc"] = i; continue
        if col["last"] is None and ("نام خانوادگ" in h_norm or "خانوادگی" in h_norm or "خانوادگي" in h_norm):
            col["last"] = i; continue
        if col["first"] is None and h_norm.strip() == "نام":
            col["first"] = i; continue
        if col["job"] is None and ("عنوان شغلی" in h_norm or "شغل" in h_norm):
            col["job"] = i; continue
        if col["contract"] is None and ("نوع قرارداد" in h_norm or "قرارداد" in h_norm):
            col["contract"] = i; continue
        if col["supply"] is None and ("تامین" in h_norm or "پخش" in h_norm or "تولید" in h_norm or "نوع تامین" in h_norm):
            col["supply"] = i; continue
        if col["deputy"] is None and "معاونت" in h_norm:
            col["deputy"] = i; continue
        if col["center"] is None and ("شبکه" in h_norm or "مرکز" in h_norm):
            col["center"] = i; continue

    if col["nc"] is None: col["nc"] = 0
    if col["first"] is None: col["first"] = 1
    if col["last"] is None: col["last"] = 2

    def to_nc(v):
        if v is None: return None
        if isinstance(v, (int, float)):
            try: return str(int(v)).zfill(10)
            except: return str(v).strip()
        s = str(v).strip().replace(".0", "").replace(" ", "")
        return s.zfill(10) if s.isdigit() else s

    def txt(v):
        if v is None: return ""
        return str(v).strip()

    ok = 0
    updated = 0
    skipped = 0
    errors = []

    con = _conn()
    for row_num, row in enumerate(ws.iter_rows(min_row=header_idx + 1, values_only=True), start=header_idx + 1):
        if not row:
            continue
        def get(idx):
            if idx is None or idx >= len(row): return None
            return row[idx]

        nc = to_nc(get(col["nc"]))
        if not nc:
            skipped += 1
            continue

        try:
            existing = con.execute("SELECT id FROM personnel WHERE national_code=?", (nc,)).fetchone()
            if existing:
                con.execute("""
                    UPDATE personnel SET
                        first_name=COALESCE(NULLIF(?, ''), first_name),
                        last_name=COALESCE(NULLIF(?, ''), last_name),
                        job_title=COALESCE(NULLIF(?, ''), job_title),
                        contract_type=COALESCE(NULLIF(?, ''), contract_type),
                        deputy=COALESCE(NULLIF(?, ''), deputy),
                        center=COALESCE(NULLIF(?, ''), center),
                        supply_type=COALESCE(NULLIF(?, ''), supply_type),
                        updated_at=datetime('now')
                    WHERE id=?
                """, (txt(get(col["first"])), txt(get(col["last"])),
                      txt(get(col["job"])), txt(get(col["contract"])),
                      txt(get(col["deputy"])), txt(get(col["center"])),
                      txt(get(col["supply"])),
                      existing["id"]))
                updated += 1
            else:
                con.execute("""
                    INSERT INTO personnel
                        (national_code, first_name, last_name, job_title,
                         contract_type, deputy, center, supply_type)
                    VALUES (?,?,?,?,?,?,?,?)
                """, (nc, txt(get(col["first"])), txt(get(col["last"])),
                      txt(get(col["job"])), txt(get(col["contract"])),
                      txt(get(col["deputy"])), txt(get(col["center"])),
                      txt(get(col["supply"]))))
                ok += 1
        except Exception as e:
            errors.append(f"ردیف {row_num}: {e}")

    con.commit()
    con.close()
    wb.close()

    return {
        "ok": ok,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:20],
        "detected_columns": col,
    }


def export_personnel():
    """خروجی اکسل کامل — همهٔ ستون‌ها + اطلاعات تجمیعی"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO

    con = _conn()

    # ==== کوئری کامل با اطلاعات تجمیعی ====
    rows = con.execute("""
        SELECT p.*,
            (SELECT t.from_deputy FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_from_deputy,
            (SELECT t.from_center FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_from_center,
            (SELECT t.to_deputy FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_to_deputy,
            (SELECT t.to_center FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_to_center,
            (SELECT t.transfer_date FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_date,
            (SELECT t.transfer_method FROM personnel_transfers t
             WHERE t.personnel_id = p.id ORDER BY t.transfer_date DESC, t.id DESC LIMIT 1) AS tr_method,
            (SELECT COUNT(*) FROM personnel_transfers t WHERE t.personnel_id = p.id) AS tr_count,
            (SELECT s.from_contract_type FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_from,
            (SELECT s.to_contract_type FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_to,
            (SELECT s.change_date FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_date,
            (SELECT s.change_method FROM personnel_status_changes s
             WHERE s.personnel_id = p.id ORDER BY s.change_date DESC, s.id DESC LIMIT 1) AS st_method,
            (SELECT COUNT(*) FROM personnel_status_changes s WHERE s.personnel_id = p.id) AS st_count
        FROM personnel p
        ORDER BY p.last_name, p.first_name
    """).fetchall()

    # ===== Sheet 1 =====
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "بانک نفرات"
    ws.sheet_view.rightToLeft = True

    headers = [
        "ردیف", "کد ملی", "نام", "نام خانوادگی",
        "عنوان شغلی", "نوع قرارداد",
        "معاونت", "شبکه/مرکز", "تامین/پخش",
        "تاریخ تبدیل وضعیت",
        "آخرین انتقال (مبدأ)", "آخرین انتقال (مقصد)", "تاریخ انتقال", "نحوه انتقال",
        "آخرین تبدیل (از)", "آخرین تبدیل (به)", "تاریخ تبدیل", "نحوه تبدیل",
    ]

    hf = Font(bold=True, color="FFFFFF", name="Tahoma", size=11)
    hfill = PatternFill("solid", start_color="1F4A7E")
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = hf
        c.fill = hfill
        c.alignment = ha

    ws.row_dimensions[1].height = 32

    for ri, r in enumerate(rows, 2):
        ws.cell(row=ri, column=1,  value=ri - 1)
        ws.cell(row=ri, column=2,  value=r["national_code"] or "")
        ws.cell(row=ri, column=3,  value=r["first_name"] or "")
        ws.cell(row=ri, column=4,  value=r["last_name"] or "")
        ws.cell(row=ri, column=5,  value=r["job_title"] or "")
        ws.cell(row=ri, column=6,  value=r["contract_type"] or "")
        ws.cell(row=ri, column=7,  value=r["deputy"] or "")
        ws.cell(row=ri, column=8,  value=r["center"] or "")

        supply = r["supply_type"] if r["deputy"] == SIMA_DEPUTY else ""
        ws.cell(row=ri, column=9,  value=supply or "")

        ws.cell(row=ri, column=10, value=r["status_changed_date"] or "")

        if r["tr_count"] and r["tr_count"] > 0:
            ws.cell(row=ri, column=11, value=r["tr_from_center"] or "")
            ws.cell(row=ri, column=12, value=r["tr_to_center"] or "")
            ws.cell(row=ri, column=13, value=r["tr_date"] or "")
            ws.cell(row=ri, column=14, value=r["tr_method"] or "")

        if r["st_count"] and r["st_count"] > 0:
            ws.cell(row=ri, column=15, value=r["st_from"] or "")
            ws.cell(row=ri, column=16, value=r["st_to"] or "")
            ws.cell(row=ri, column=17, value=r["st_date"] or "")
            ws.cell(row=ri, column=18, value=r["st_method"] or "")

    widths = [6, 14, 12, 16, 18, 16, 18, 20, 12, 14, 20, 20, 14, 14, 14, 14, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet 2 =====
    ws2 = wb.create_sheet("سابقه انتقالات")
    ws2.sheet_view.rightToLeft = True

    h2 = ["ردیف", "کد ملی", "نام", "نام خانوادگی", "از مرکز", "به مرکز",
          "تاریخ", "نحوه", "توضیحات"]
    for i, h in enumerate(h2, 1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = hf; c.fill = hfill; c.alignment = ha

    transfers = con.execute("""
        SELECT t.*, p.national_code, p.first_name, p.last_name
        FROM personnel_transfers t
        JOIN personnel p ON p.id = t.personnel_id
        ORDER BY t.transfer_date DESC, t.id DESC
    """).fetchall()

    for ri, t in enumerate(transfers, 2):
        ws2.cell(row=ri, column=1, value=ri - 1)
        ws2.cell(row=ri, column=2, value=t["national_code"] or "")
        ws2.cell(row=ri, column=3, value=t["first_name"] or "")
        ws2.cell(row=ri, column=4, value=t["last_name"] or "")
        ws2.cell(row=ri, column=5, value=(t["from_deputy"] or "") + " / " + (t["from_center"] or ""))
        ws2.cell(row=ri, column=6, value=(t["to_deputy"] or "") + " / " + (t["to_center"] or ""))
        ws2.cell(row=ri, column=7, value=t["transfer_date"] or "")
        ws2.cell(row=ri, column=8, value=t["transfer_method"] or "")
        ws2.cell(row=ri, column=9, value=t["note"] or "")

    for i, w in enumerate([6, 14, 12, 16, 30, 30, 14, 14, 30], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet 3 =====
    ws3 = wb.create_sheet("سابقه تبدیل وضعیت")
    ws3.sheet_view.rightToLeft = True

    h3 = ["ردیف", "کد ملی", "نام", "نام خانوادگی", "از وضعیت", "به وضعیت",
          "تاریخ", "نحوه", "توضیحات"]
    for i, h in enumerate(h3, 1):
        c = ws3.cell(row=1, column=i, value=h)
        c.font = hf; c.fill = hfill; c.alignment = ha

    changes = con.execute("""
        SELECT s.*, p.national_code, p.first_name, p.last_name
        FROM personnel_status_changes s
        JOIN personnel p ON p.id = s.personnel_id
        ORDER BY s.change_date DESC, s.id DESC
    """).fetchall()

    for ri, s in enumerate(changes, 2):
        ws3.cell(row=ri, column=1, value=ri - 1)
        ws3.cell(row=ri, column=2, value=s["national_code"] or "")
        ws3.cell(row=ri, column=3, value=s["first_name"] or "")
        ws3.cell(row=ri, column=4, value=s["last_name"] or "")
        ws3.cell(row=ri, column=5, value=s["from_contract_type"] or "")
        ws3.cell(row=ri, column=6, value=s["to_contract_type"] or "")
        ws3.cell(row=ri, column=7, value=s["change_date"] or "")
        ws3.cell(row=ri, column=8, value=s["change_method"] or "")
        ws3.cell(row=ri, column=9, value=s["note"] or "")

    for i, w in enumerate([6, 14, 12, 16, 18, 18, 14, 14, 30], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    con.close()

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
