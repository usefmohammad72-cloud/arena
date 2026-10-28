# -*- coding: utf-8 -*-
"""الگو مصوب — دیتابیس و منطق"""
import sqlite3, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
from app_paths import get_db_path

SIMA_DEPUTY = "معاونت سیما"


def _conn():
    c = sqlite3.connect(str(get_db_path()), timeout=30)
    c.row_factory = sqlite3.Row
    return c


def init_templates_table():
    con = _conn()
    con.executescript("""
        CREATE TABLE IF NOT EXISTS personnel_templates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            deputy TEXT NOT NULL,
            center TEXT NOT NULL,
            supply_type TEXT DEFAULT '',
            job_title TEXT DEFAULT '',
            count INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now')),
            updated_at TEXT DEFAULT (datetime('now')),
            UNIQUE(deputy, center, supply_type, job_title)
        );
        CREATE INDEX IF NOT EXISTS idx_templates_dcs
            ON personnel_templates(deputy, center, supply_type);

        CREATE TABLE IF NOT EXISTS personnel_shifts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            deputy TEXT NOT NULL,
            center TEXT NOT NULL,
            supply_type TEXT DEFAULT '',
            job_title TEXT DEFAULT '',
            employment_type TEXT NOT NULL,
            shift_amount REAL DEFAULT 0,
            shift_count INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now')),
            updated_at TEXT DEFAULT (datetime('now')),
            UNIQUE(deputy, center, supply_type, job_title, employment_type)
        );
        CREATE INDEX IF NOT EXISTS idx_shifts_dcs
            ON personnel_shifts(deputy, center, supply_type);
    """)
    con.commit()
    con.close()


def get_summary():
    """جدول اصلی الگو مصوب — با شمارش نفرات هر قرارداد"""
    con = _conn()

    # جمع الگو مصوب per (deputy, center, supply)
    template_sums = {}
    for r in con.execute("""
        SELECT deputy, center, COALESCE(supply_type,'') AS supply,
               SUM(count) AS total
        FROM personnel_templates
        GROUP BY deputy, center, supply
    """).fetchall():
        template_sums[(r["deputy"], r["center"], r["supply"])] = r["total"] or 0

    # شمارش نفرات per (deputy, center, supply, contract_type)
    personnel_counts = {}
    for r in con.execute("""
        SELECT deputy, center, COALESCE(supply_type,'') AS supply,
               contract_type, COUNT(*) AS cnt
        FROM personnel
        GROUP BY deputy, center, supply, contract_type
    """).fetchall():
        key = (r["deputy"], r["center"], r["supply"])
        personnel_counts.setdefault(key, {})[r["contract_type"] or ""] = r["cnt"]

    all_keys = set(template_sums.keys()) | set(personnel_counts.keys())

    rows = []
    for key in sorted(all_keys):
        deputy, center, supply = key
        counts = personnel_counts.get(key, {})

        formal = counts.get("رسمی", 0)
        contractual = counts.get("پیمانی", 0)
        buy_services = counts.get("قرارداد خرید خدمات", 0)
        corporate = counts.get("شرکتی", 0)
        article28 = counts.get("ماده 28", 0)
        total_formal = formal + contractual + buy_services + corporate

        rows.append({
            "deputy": deputy,
            "center": center,
            "supply_type": supply if deputy == SIMA_DEPUTY else "",
            "template_count": template_sums.get(key, 0),
            "formal": formal,
            "contractual": contractual,
            "buy_services": buy_services,
            "corporate": corporate,
            "article28": article28,
            "total_formal": total_formal,
        })

    con.close()
    return rows


def get_details(deputy, center, supply_type):
    """جزئیات الگو مصوب per job_title"""
    from web.personnel import JOB_TITLES
    con = _conn()

    templates = {}
    for r in con.execute("""
        SELECT job_title, count FROM personnel_templates
        WHERE deputy=? AND center=? AND COALESCE(supply_type,'')=?
    """, (deputy, center, supply_type or "")).fetchall():
        templates[r["job_title"]] = r["count"]

    personnel = {}
    for r in con.execute("""
        SELECT job_title, COUNT(*) AS cnt
        FROM personnel
        WHERE deputy=? AND center=? AND COALESCE(supply_type,'')=?
        GROUP BY job_title
    """, (deputy, center, supply_type or "")).fetchall():
        personnel[r["job_title"]] = r["cnt"]

    result = []
    for jt in JOB_TITLES:
        result.append({
            "job_title": jt,
            "template": templates.get(jt, 0),
            "actual": personnel.get(jt, 0),
        })

    con.close()
    return result


def save_batch(deputy, center, supply_type, items):
    """ذخیرهٔ گروهی الگو مصوب — items: [{job_title, count}, ...]"""
    con = _conn()
    saved = 0
    for item in items:
        jt = (item.get("job_title") or "").strip()
        try:
            cnt = int(item.get("count") or 0)
        except (TypeError, ValueError):
            cnt = 0
        if not jt:
            continue
        con.execute("""
            INSERT INTO personnel_templates (deputy, center, supply_type, job_title, count)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(deputy, center, supply_type, job_title) DO UPDATE SET
                count = excluded.count,
                updated_at = datetime('now')
        """, (deputy, center, supply_type or "", jt, cnt))
        saved += 1
    con.commit()
    con.close()
    return saved


def import_excel(xlsx_path):
    """ایمپورت الگو مصوب از اکسل — معاونت | شبکه/مرکز | تامین/پخش | عنوان شغلی | تعداد"""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_idx = 1
    header = None
    for ri in range(1, 6):
        try:
            row = next(ws.iter_rows(min_row=ri, max_row=ri, values_only=True))
        except StopIteration:
            break
        if not row:
            continue
        text = " ".join(str(v) for v in row if v is not None)
        if "معاونت" in text or "شبکه" in text:
            header_idx = ri
            header = row
            break

    if header is None:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    col = {"deputy": None, "center": None, "supply": None, "job": None, "count": None}
    for i, h in enumerate(header):
        if h is None:
            continue
        h_norm = str(h).strip().replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
        if col["deputy"] is None and "معاونت" in h_norm:
            col["deputy"] = i; continue
        if col["center"] is None and ("شبکه" in h_norm or "مرکز" in h_norm):
            col["center"] = i; continue
        if col["supply"] is None and ("تامین" in h_norm or "پخش" in h_norm or "تولید" in h_norm):
            col["supply"] = i; continue
        if col["job"] is None and ("عنوان شغلی" in h_norm or "شغل" in h_norm):
            col["job"] = i; continue
        if col["count"] is None and ("تعداد" in h_norm or "الگو" in h_norm or "مصوب" in h_norm):
            col["count"] = i; continue

    if col["deputy"] is None: col["deputy"] = 0
    if col["center"] is None: col["center"] = 1
    if col["supply"] is None: col["supply"] = 2
    if col["job"] is None: col["job"] = 3
    if col["count"] is None: col["count"] = 4

    def txt(v):
        if v is None: return ""
        return str(v).strip()

    ok = 0
    skipped = 0
    errors = []

    con = _conn()
    for row_num, row in enumerate(
        ws.iter_rows(min_row=header_idx + 1, values_only=True),
        start=header_idx + 1
    ):
        if not row:
            continue

        def get(idx):
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        deputy = txt(get(col["deputy"]))
        center = txt(get(col["center"]))
        supply = txt(get(col["supply"]))
        job = txt(get(col["job"]))
        try:
            count = int(float(get(col["count"]) or 0))
        except (ValueError, TypeError):
            count = 0

        if not deputy or not center or not job:
            skipped += 1
            continue

        if deputy != SIMA_DEPUTY:
            supply = ""

        try:
            con.execute("""
                INSERT INTO personnel_templates
                    (deputy, center, supply_type, job_title, count)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(deputy, center, supply_type, job_title) DO UPDATE SET
                    count = excluded.count,
                    updated_at = datetime('now')
            """, (deputy, center, supply, job, count))
            ok += 1
        except Exception as e:
            errors.append(f"ردیف {row_num}: {e}")

    con.commit()
    con.close()
    wb.close()

    return {"ok": ok, "skipped": skipped, "errors": errors[:20]}


def get_full_matrix(deputy_filter=None, center_filter=None, supply_filter=None):
    """
    Cross join کامل:
    همهٔ معاونت‌ها × همهٔ مراکز × تامین/پخش (برای سیما) × همهٔ ۱۸ عنوان شغلی
    مستقل از اینکه داده داشته باشن یا نه.
    """
    from web.personnel import JOB_TITLES

    con = _conn()

    # ۱) همهٔ ترکیب‌های (معاونت، مرکز، تامین/پخش)
    combos = set()
    for r in con.execute("""
        SELECT DISTINCT deputy, center_name
        FROM centers_structure
        WHERE deputy IS NOT NULL AND deputy != ''
          AND center_name IS NOT NULL AND center_name != ''
    """).fetchall():
        dep = r["deputy"]
        cen = r["center_name"]
        if dep == SIMA_DEPUTY:
            # برای سیما: همیشه ۳ ردیف جدا
            for s in ("تامین", "پخش", "تولید"):
                combos.add((dep, cen, s))
        else:
            combos.add((dep, cen, ""))

    # ۲) شمارش نفرات per (deputy, center, supply, job_title, contract_type)
    counts_map = {}
    for r in con.execute("""
        SELECT deputy, center, COALESCE(supply_type,'') AS supply,
               job_title, contract_type, COUNT(*) AS cnt
        FROM personnel
        GROUP BY deputy, center, supply, job_title, contract_type
    """).fetchall():
        key = (r["deputy"], r["center"], r["supply"], r["job_title"] or "")
        counts_map.setdefault(key, {})[r["contract_type"] or ""] = r["cnt"]

    # ۳) الگو مصوب
    template_map = {}
    for r in con.execute("""
        SELECT deputy, center, COALESCE(supply_type,'') AS supply,
               job_title, count
        FROM personnel_templates
    """).fetchall():
        key = (r["deputy"], r["center"], r["supply"], r["job_title"] or "")
        template_map[key] = r["count"] or 0

    # ۴) ساخت ردیف‌ها
    rows = []
    for (dep, cen, sup) in sorted(combos):
        if deputy_filter and deputy_filter != "همه" and dep != deputy_filter:
            continue
        if center_filter and center_filter != "همه" and cen != center_filter:
            continue
        if supply_filter and supply_filter != "همه":
            # فیلتر تامین/پخش فقط برای سیما اعمال می‌شه
            if dep == SIMA_DEPUTY and sup != supply_filter:
                continue

        for jt in JOB_TITLES:
            key = (dep, cen, sup, jt)
            counts = counts_map.get(key, {})
            formal = counts.get("رسمی", 0)
            contractual = counts.get("پیمانی", 0)
            buy = counts.get("قرارداد خرید خدمات", 0)
            corp = counts.get("شرکتی", 0)
            art28 = counts.get("ماده 28", 0)

            rows.append({
                "deputy": dep,
                "center": cen,
                "supply_type": sup if dep == SIMA_DEPUTY else "",
                "job_title": jt,
                "template": template_map.get(key, 0),
                "formal": formal,
                "contractual": contractual,
                "buy_services": buy,
                "corporate": corp,
                "article28": art28,
                "total_formal": formal + contractual + buy + corp,
            })

    con.close()
    return rows


def save_matrix(rows):
    """ذخیرهٔ ماتریس — فقط ردیف‌های تغییر یافته"""
    con = _conn()
    upsert_data = []
    delete_keys = []

    for row in rows:
        deputy = (row.get("deputy") or "").strip()
        center = (row.get("center") or "").strip()
        supply = (row.get("supply_type") or "").strip()
        jt = (row.get("job_title") or "").strip()
        try:
            cnt = int(row.get("count") or 0)
        except (ValueError, TypeError):
            cnt = 0

        if not deputy or not center or not jt:
            continue
        if deputy != SIMA_DEPUTY:
            supply = ""

        if cnt <= 0:
            delete_keys.append((deputy, center, supply, jt))
        else:
            upsert_data.append((deputy, center, supply, jt, cnt))

    # حذف صفرها
    if delete_keys:
        con.executemany("""
            DELETE FROM personnel_templates
            WHERE deputy=? AND center=? AND COALESCE(supply_type,'')=? AND job_title=?
        """, delete_keys)

    # upsert
    if upsert_data:
        con.executemany("""
            INSERT INTO personnel_templates
                (deputy, center, supply_type, job_title, count)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(deputy, center, supply_type, job_title) DO UPDATE SET
                count = excluded.count,
                updated_at = datetime('now')
        """, upsert_data)

    con.commit()
    con.close()
    return len(upsert_data)

def get_shifts_matrix(deputy_filter=None, center_filter=None, supply_filter=None):
    """ماتریس شیفت‌ها — cross join کامل: معاونت × شبکه × تامین/پخش × عنوان شغلی × نوع استخدام"""
    from web.personnel import JOB_TITLES

    con = _conn()

    # ۱) ترکیب‌های (معاونت، شبکه، تامین/پخش)
    combos = set()
    for r in con.execute("""
        SELECT DISTINCT deputy, center_name FROM centers_structure
        WHERE deputy IS NOT NULL AND deputy != ''
          AND center_name IS NOT NULL AND center_name != ''
    """).fetchall():
        dep = r["deputy"]
        cen = r["center_name"]
        if dep == SIMA_DEPUTY:
            for s in ("تامین", "پخش", "تولید"):
                combos.add((dep, cen, s))
        else:
            combos.add((dep, cen, ""))

    # ۲) خواندن شیفت‌های ذخیره‌شده
    shifts_map = {}
    for r in con.execute("""
        SELECT deputy, center, COALESCE(supply_type,'') AS supply,
               job_title, employment_type, shift_amount, shift_count
        FROM personnel_shifts
    """).fetchall():
        key = (r["deputy"], r["center"], r["supply"], r["job_title"] or "", r["employment_type"])
        shifts_map[key] = {
            "amount": r["shift_amount"] or 0,
            "count": r["shift_count"] or 0,
        }

    # ۳) ساخت ماتریس
    rows = []
    for (dep, cen, sup) in sorted(combos):
        if deputy_filter and deputy_filter != "همه" and dep != deputy_filter:
            continue
        if center_filter and center_filter != "همه" and cen != center_filter:
            continue
        if supply_filter and supply_filter != "همه" and dep == SIMA_DEPUTY and sup != supply_filter:
            continue

        for jt in JOB_TITLES:
            for emp_type in ("رسمی", "تبصره28"):
                key = (dep, cen, sup, jt, emp_type)
                data = shifts_map.get(key, {"amount": 0, "count": 0})
                rows.append({
                    "deputy": dep,
                    "center": cen,
                    "supply_type": sup if dep == SIMA_DEPUTY else "",
                    "job_title": jt,
                    "employment_type": emp_type,
                    "shift_amount": data["amount"],
                    "shift_count": data["count"],
                })

    con.close()
    return rows


def save_shift(deputy, center, supply_type, job_title, employment_type, amount, count):
    """ذخیرهٔ یک ردیف شیفت"""
    con = _conn()
    if deputy != SIMA_DEPUTY:
        supply_type = ""

    con.execute("""
        INSERT INTO personnel_shifts
            (deputy, center, supply_type, job_title, employment_type, shift_amount, shift_count)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(deputy, center, supply_type, job_title, employment_type) DO UPDATE SET
            shift_amount = excluded.shift_amount,
            shift_count = excluded.shift_count,
            updated_at = datetime('now')
    """, (deputy, center, supply_type, job_title, employment_type,
          float(amount or 0), int(count or 0)))
    con.commit()
    con.close()
    return True

